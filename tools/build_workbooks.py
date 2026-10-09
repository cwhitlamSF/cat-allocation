"""
Builds the CCA workbooks from the layouts in cca_tables.py:

  CCA_ConfigSpecs.xlsx  spec definition tables read by config_creator -> CCA_config.ini
  CCA_Template.xlsm     the input workbook users fill in, with dropdowns, the result sheets and the VBA
  CCA_Example.xlsm      (--example) the template filled in for the sample data in example_data/

The spec workbook and template are always generated together, so they can't drift apart.
Rerun after changing cca_tables.py:

    python tools/build_workbooks.py [--example] [--vba vba/vbaProject.bin | --novba]

The VBA project (the Run macro and the xlwings module) is copied in from vba/vbaProject.bin, so a
rebuild keeps it. To change the VBA: edit it in Excel in CCA_Template.xlsm, then save the project back
with  python tools/build_workbooks.py --save-vba CCA_Template.xlsm  (copies its vbaProject.bin to vba/).
--novba writes .xlsx workbooks without VBA.
"""
import os
import sys
import tempfile
import zipfile

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
from cca_tables import LISTS, TABLES, RESULT_TABLES      # noqa: E402

VBA_PROJECT = os.path.join(ROOT, 'vba', 'vbaProject.bin')

VALIDATION_ROWS = 300
NOTE_FONT = Font(italic=True, color='808080', size=9)
TITLE_FONT = Font(size=16, bold=True, color='1D5AA5')


def _add_table(ws, name, header, rows, first_row, style='TableStyleMedium2'):
    """Excel table with its header at first_row; an empty table keeps one blank data row."""
    for j, h in enumerate(header, 1):
        ws.cell(first_row, j, h)
    for i, r in enumerate(rows, 1):
        for j, v in enumerate(r, 1):
            ws.cell(first_row + i, j, v)
    last = first_row + max(len(rows), 1)
    t = Table(displayName=name, ref=f"A{first_row}:{get_column_letter(len(header))}{last}")
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(t)
    return t


def _value(v):
    return {'TRUE': True, 'FALSE': False}.get(v, v) if isinstance(v, str) else v


HOME_LINES = [
    'Allocates the upfront premium of the upper cat layers to individual policies, from RMS ELTs, without simulation.',
    '',
    '1. General: the output folder and where the event-grouped policy ELT files are.',
    '2. Prepare Data (once per RMS run): where the raw policy ELT is; set Run Prepare Data to TRUE to build the',
    '   event-grouped files on the next Run. Set it back to FALSE afterwards.',
    '3. Specify Runs: one row per run, each with a Run Name and its method settings (distortion, basis, sharing,',
    '   split, ...). One Run does every row with Include = TRUE; the data is read once for all of them.',
    '4. Data Sources: the BU x region ELTs (gross and net of per-risk), the exposure data and the policy-to-LOB map.',
    '5. Subjects: named sets of LOBs and regions that the layers apply to.',
    '6. Inuring Layers: lower towers and the FHCF proxy, applied in Stage order before the upper layers.',
    '7. Upper Layers: the layers whose premium is allocated (100% premium and placement; aggregate terms if used).',
    '',
    'Results: Run Summary (one row per run), Layer Summary, LOB Summary, Diagnostics and Benchmark, each with a',
    'Run Name column; policy-level results in <Output Folder>/<Run Name>/, and all runs side by side in',
    'premium_by_policy_all_runs.parquet. Files written are listed on the Output Log sheet.',
    'Problems found before or during the run are listed on the Data Issues sheet.',
    'Relative paths are relative to this workbook\'s folder.',
]


def build_template(path, example=None, vba=None):
    """Write the template (or, with `example`, a filled-in copy). With `vba` (a vbaProject.bin),
    the workbook is saved as .xlsm with that VBA project."""
    example = example or {}
    wb = openpyxl.Workbook()
    home = wb.active
    home.title = 'Home'
    home['A1'] = 'Cat Cost Allocation'
    home['A1'].font = TITLE_FONT
    for i, line in enumerate(HOME_LINES, 3):
        home.cell(i, 1, line)
    home.column_dimensions['A'].width = 120

    lists = wb.create_sheet('Lists')
    for j, (name, values) in enumerate(LISTS.items(), 1):
        lists.cell(1, j, name)
        for i, v in enumerate(values, 2):
            lists.cell(i, j, v)
        col = get_column_letter(j)
        wb.defined_names[name] = DefinedName(name, attr_text=f"Lists!${col}$2:${col}${len(values) + 1}")
    lists.sheet_state = 'hidden'

    for spec in TABLES:
        ws = wb.create_sheet(spec['sheet'])
        header = [c[0] for c in spec['columns']]
        if spec.get('info'):
            values = example.get(spec['table'], {})
            rows = [[name, _value(values.get(name, default)), note]
                    for name, default, lst, note in spec['settings']]
            _add_table(ws, spec['table'], header, rows, first_row=1)
            for i, (name, default, lst, note) in enumerate(spec['settings'], 2):
                if lst:
                    dv = DataValidation(type='list', formula1=f'={lst}', allow_blank=True)
                    dv.error = f'Choose a value from the list for {name}'
                    dv.add(f"B{i}")
                    ws.add_data_validation(dv)
                ws.cell(i, 3).font = NOTE_FONT
            for j, c in enumerate(spec['columns'], 1):
                ws.column_dimensions[get_column_letter(j)].width = c[5]
            ws.freeze_panes = 'A2'
            continue
        rows = [[_value(v) for v in r] + [None] * (len(header) - len(r)) for r in example.get(spec['table'], [])]
        _add_table(ws, spec['table'], header, rows, first_row=2)
        for j, (col, _type, default, required, lst, width, note) in enumerate(spec['columns'], 1):
            letter = get_column_letter(j)
            ws.column_dimensions[letter].width = width
            text = ' '.join(x for x in [note, '(required)' if required else '',
                                        f'Default: {default}.' if default is not None else ''] if x)
            cell = ws.cell(1, j, text)
            cell.font = NOTE_FONT
            cell.alignment = Alignment(wrap_text=True, vertical='bottom')
            if lst:
                dv = DataValidation(type='list', formula1=f'={lst}', allow_blank=True)
                dv.error = f'Choose a value from the list for {col}'
                dv.add(f"{letter}3:{letter}{2 + VALIDATION_ROWS}")
                ws.add_data_validation(dv)
        ws.row_dimensions[1].height = 60
        ws.freeze_panes = 'A3'

    for res in RESULT_TABLES:
        ws = wb.create_sheet(res['sheet'])
        _add_table(ws, res['table'], [c[0] for c in res['columns']], [], first_row=1, style='TableStyleLight9')
        for j, (c, width) in enumerate(res['columns'], 1):
            ws.column_dimensions[get_column_letter(j)].width = width
        ws.freeze_panes = 'A2'

    # Data Issues: names written by modelkit.status.ExcelStatus
    di = wb.create_sheet('Data Issues')
    di['A1'], di['B1'] = 'Run status', None
    di['A3'] = 'Data warnings from the last run'
    di['A1'].font = di['A3'].font = Font(bold=True)
    di.column_dimensions['A'].width = 120
    wb.defined_names['runstatus'] = DefinedName('runstatus', attr_text="'Data Issues'!$B$1")
    wb.defined_names['startdatawarnings'] = DefinedName('startdatawarnings', attr_text="'Data Issues'!$A$4")

    # Model Path: where the Run macro finds the exe (same names as RSA and DTX)
    mp = wb.create_sheet('Model Path')
    mp['A1'], mp['B1'] = 'Executable path', r'C:\CCA\CCA.exe'
    mp['A2'] = 'Path to the CCA exe used by the Run button. Change it if the tool is installed elsewhere.'
    mp['A2'].font = NOTE_FONT
    mp.column_dimensions['A'].width = 18
    mp.column_dimensions['B'].width = 60
    wb.defined_names['_executablepath'] = DefinedName('_executablepath', attr_text="'Model Path'!$B$1")
    if vba:
        _attach_vba(wb, vba)
    wb.save(path)


def _attach_vba(wb, vba_bin):
    """Give the workbook a VBA project. openpyxl writes a macro-enabled workbook (content types and
    the vbaProject relationship) when vba_archive is set; the document modules in the project are
    matched to sheets by code name, so the sheets get ThisWorkbook / Sheet1, Sheet2, ... as far as the
    project has them (Excel adds modules for any further sheets)."""
    tmp = tempfile.NamedTemporaryFile(suffix='.zip', delete=False)
    tmp.close()
    with zipfile.ZipFile(tmp.name, 'w') as z:
        z.write(vba_bin, 'xl/vbaProject.bin')
        # openpyxl also reads these two parts from the archive (to carry over custom UI and
        # content types); standard minimal versions
        z.writestr('[Content_Types].xml',
                   '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/>'
                   '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                   '<Default Extension="xml" ContentType="application/xml"/></Types>')
        z.writestr('_rels/.rels',
                   '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                   'relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
    n_modules = _document_modules(vba_bin)
    wb.vba_archive = zipfile.ZipFile(tmp.name)
    wb.code_name = 'ThisWorkbook'
    for i, ws in enumerate(wb.worksheets, 1):
        if i <= n_modules:
            ws.sheet_properties.codeName = f'Sheet{i}'


def _document_modules(vba_bin):
    """Number of SheetN document modules in a vbaProject.bin (by stream name)."""
    try:
        import olefile
        ole = olefile.OleFileIO(vba_bin)
        names = {e[-1] for e in ole.listdir() if len(e) == 2 and e[0] == 'VBA'}
        n = 0
        while f'Sheet{n + 1}' in names:
            n += 1
        return n
    except Exception:
        return 0


def build_spec_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Config Specs'
    definition = [
        ['Model Spec Sources', None, 'Tbl_dataInputSpecs', 'Information',
         'Sheet Name,Table Name,Key Column (s),Default Source,Flag Info Table,Table Has Headers', None],
        ['Spec Data Formats', None, 'Tbl_specDataFormats', 'Spec Table,Column',
         'Data Type,Default Value,Exclude if Null', None],
    ]
    _add_table(ws, 'Tbl_specDefinition', ['Spec Key', 'Panel', 'Source Table', 'Key Column (s)',
                                          'Value Column (s)', 'Comments'], definition, first_row=1)
    inputs = [[t['spec'], t['sheet'], t['table'], t['keys'] or None, None, bool(t.get('info')), True] for t in TABLES]
    ws2 = wb.create_sheet('Inputs')
    _add_table(ws2, 'Tbl_dataInputSpecs', ['Information', 'Sheet Name', 'Table Name', 'Key Column (s)',
                                           'Default Source', 'Flag Info Table', 'Table Has Headers'], inputs, first_row=1)
    formats = [[t['spec'], c[0], c[1], _value(c[2]), True if c[3] else None] for t in TABLES for c in t['columns']]
    ws3 = wb.create_sheet('Formats')
    _add_table(ws3, 'Tbl_specDataFormats', ['Spec Table', 'Column', 'Data Type', 'Default Value', 'Exclude if Null'],
               formats, first_row=1)
    for sheet in (ws, ws2, ws3):
        for col in 'ABCDEFG':
            sheet.column_dimensions[col].width = 22
    wb.save(path)


# The worked example: the synthetic book written by tools/make_example_data.py into example_data/
EXAMPLE = {
    'GeneralSpecs': {'Output Folder': 'example_data/output', 'Policy ELT Folder': 'example_data/ev_buckets',
                     'Notes': 'Built by tools/build_workbooks.py --example'},
    'PrepareSpecs': {'Run Prepare Data': 'TRUE', 'Source Type': 'Parquet Files', 'Input Folder': 'example_data/raw_policy_elt',
                     'Input File Pattern': 'chunk_*.parquet', 'Target Rows Per File': '20000',
                     'Column Renames': 'Loss=PERSPVALUE'},
    'SpecifyRuns': [
        # Run Name, Include, Distortion, Basis, Inuring Sharing, Policy Split, Bounded, Clamp, RP, Benchmark, Grid, Comments
        ['Base', 'TRUE', 'wang', 'occurrence', 'conditional', 'conditional', 'TRUE', 'FALSE', 'TRUE', 'TRUE', 500, 'Recommended method'],
        ['PH sensitivity', 'TRUE', 'ph', 'occurrence', 'conditional', 'conditional', 'TRUE', 'FALSE', 'TRUE', 'FALSE', 500, None],
        ['Aggregate basis', 'TRUE', 'wang', 'aggregate', 'conditional', 'conditional', 'TRUE', 'FALSE', 'TRUE', 'FALSE', 500, None],
        ['Pro rata sharing', 'TRUE', 'wang', 'occurrence', 'prorata', 'conditional', 'TRUE', 'FALSE', 'TRUE', 'FALSE', 500, None],
        ['Mean split', 'TRUE', 'wang', 'occurrence', 'conditional', 'mean', 'TRUE', 'FALSE', 'TRUE', 'FALSE', 500, 'Same lookups as Base'],
        ['Not run', 'FALSE', 'identity', None, None, None, None, None, None, None, None, 'Include = FALSE'],
    ],
    'DataSources': [
        ['BU ELT Gross', 'Parquet', 'example_data/bu_elt_gross.parquet'],
        ['BU ELT Net', 'Parquet', 'example_data/bu_elt_net.parquet'],
        ['Exposure', 'CSV', 'example_data/exposure.csv', None, None, None, None, None, 'PolicyID=POLICYID'],
        ['Policy LOB', 'CSV', 'example_data/exposure.csv', None, None, None, None, None, 'PolicyID=POLICYID'],
    ],
    'Subjects': [
        ['ALL', None, None, 'Every LOB, every region'],
        ['HO', 'HO', None, None],
        ['CO', 'CO', None, None],
        ['FHCF', 'HO', 'US_FL', 'FHCF proxy: HO in Florida'],
    ],
    'InuringLayers': [
        [1, 'FHCF', 5e6, 20e6, 0.9, 'FHCF', 'FALSE', 'TRUE'],
        [2, 'HO 1', 3e6, 5e6, 1.0, 'HO', None, 'TRUE'],
        [2, 'HO 2', 8e6, 10e6, 0.95, 'HO', None, 'TRUE'],
        [2, 'CO 1', 4e6, 8e6, 1.0, 'CO', None, 'TRUE'],
    ],
    'UpperLayers': [
        ['U1', 10e6, 15e6, 'ALL', 5e6, 0.6, None, 1],
        ['U2', 25e6, 25e6, 'ALL', 1.5e6, 1.0, None, 1],
        ['U3 agg', 4e6, 6e6, 'HO', 1.2e6, 0.75, 'aggregate', 1],
    ],
}


def save_vba(xlsm):
    """Copy the VBA project out of an .xlsm into vba/vbaProject.bin (after editing the VBA in Excel)."""
    with zipfile.ZipFile(xlsm) as z:
        data = z.read('xl/vbaProject.bin')
    with open(VBA_PROJECT, 'wb') as f:
        f.write(data)
    return VBA_PROJECT


if __name__ == '__main__':
    args = sys.argv[1:]
    if '--save-vba' in args:
        print('Wrote', save_vba(args[args.index('--save-vba') + 1]))
        sys.exit(0)
    vba = None if '--novba' in args else (args[args.index('--vba') + 1] if '--vba' in args else VBA_PROJECT)
    if vba and not os.path.exists(vba):
        sys.exit(f"No VBA project at {vba}: pass --vba <vbaProject.bin> or --novba")
    ext = '.xlsm' if vba else '.xlsx'
    build_spec_workbook(os.path.join(ROOT, 'CCA_ConfigSpecs.xlsx'))
    build_template(os.path.join(ROOT, 'CCA_Template' + ext), vba=vba)
    print(f'Wrote CCA_ConfigSpecs.xlsx and CCA_Template{ext}')
    if '--example' in args:
        build_template(os.path.join(ROOT, 'CCA_Example' + ext), EXAMPLE, vba=vba)
        print(f'Wrote CCA_Example{ext}')
